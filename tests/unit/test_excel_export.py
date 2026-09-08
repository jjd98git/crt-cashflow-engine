"""Per-tranche allocations in the run API and the values workbook export.

What these prove: every tranche balance moves only by what the engine allocated to it
(before - principal - write-down + write-up + increase == after) on both example
scenarios; the workbook reopens with the expected sheets, its Structure sheet reproduces
the spec 02 section 12 balances after Payment Date 1, the Charts sheet holds the ten
charts, and the export is deterministic (same result -> same bytes; same scenario, other
timestamp -> same cell values except the README timestamp)."""

from __future__ import annotations

import csv
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

from crt.api import RunResult, export_csv, run_scenario
from crt.excel.__main__ import main as excel_main
from crt.excel.structure_workbook import (
    BOTTOM_UP_ORDER,
    SHEET_NAMES,
    STRUCTURE_HEADERS,
    VALUES_ONLY_SENTENCE,
    build_structure_workbook,
    export_structure_workbook,
    structure_workbook_bytes,
)
from crt.io.deal_terms import NOTE_CLASSES, TRANCHE_ORDER
from crt.scenarios.loader import load_scenario

pytestmark = pytest.mark.fast

D = Decimal
EXPECTED_CHART_COUNT = 10  # (a) (b) (c) (d) + four Note charts (e) + two line charts (f)


def _run(project_root: Path, file_name: str, timestamp: str) -> RunResult:
    source = project_root / "scenarios" / file_name
    return run_scenario(
        load_scenario(source),
        project_root=project_root,
        scenario_name=source.stem,
        scenario_source=source,
        timestamp=timestamp,
    )


@pytest.fixture(scope="module")
def pricing(project_root: Path) -> RunResult:
    return _run(project_root, "pricing_speed.yaml", "2026-01-01T00:00:00Z")


@pytest.fixture(scope="module")
def stress(project_root: Path) -> RunResult:
    return _run(project_root, "stress_5cpr_2_5cer.yaml", "2026-01-01T00:00:00Z")


# --------------------------------------------------------------------------------------
# API: allocations
# --------------------------------------------------------------------------------------


@pytest.mark.parametrize("which", ["pricing", "stress"])
def test_allocation_identity_per_tranche(which: str, request: pytest.FixtureRequest) -> None:
    result: RunResult = request.getfixturevalue(which)
    previous_after: dict[str, Decimal] | None = None
    for row in result.structure:
        for tranche in TRANCHE_ORDER:
            assert (
                row.balances_before[tranche]
                - row.principal_allocated[tranche]
                - row.write_down_allocated[tranche]
                + row.write_up_allocated[tranche]
                + row.increase_allocated[tranche]
                == row.balances_after[tranche]
            ), (row.payment_date_number, tranche)
            if tranche != "A-H":
                assert row.increase_allocated[tranche] == D(0)
        if previous_after is not None:
            assert row.balances_before == previous_after
        previous_after = row.balances_after
        # Before the Maturity Date the Reference Tranches sum to the pool UPB (A-H absorbs
        # every difference).  On the Maturity Date the tranches are paid to zero (spec 02
        # section 9) while the pool still carries its UPB, so the identity ends there.
        if not row.is_maturity_date:
            assert sum(row.balances_after.values(), D(0)) == row.pool_upb_end
        # The Notes' principal is exactly the tranche principal of their tranche.
        for note in NOTE_CLASSES:
            cashflow = next(
                c
                for c in result.note_cashflows
                if c.note == note and c.payment_date_number == row.payment_date_number
            )
            assert cashflow.principal_paid == row.principal_allocated[note]
            assert cashflow.write_down == row.write_down_allocated[note]
    last = result.structure[-1]
    assert last.is_maturity_date
    for tranche in TRANCHE_ORDER:
        assert last.balances_after[tranche] == D(0)
        summary = result.tranche_summary_for(tranche)
        assert summary.total_write_downs == last.cumulative_write_down[tranche]
        assert summary.final_balance == D(0)
        assert (
            summary.initial_class_notional_amount
            - summary.total_principal
            - summary.total_write_downs
            + summary.total_write_ups
            + summary.total_increases
            == summary.final_balance
        )


def test_pricing_speed_payment_date_1_allocations_match_spec_02_section_12(
    pricing: RunResult,
) -> None:
    first = pricing.structure[0]
    assert first.principal_allocated["A-1"] == D("10346250.00")
    assert first.principal_allocated["A-1H"] == D("545988.08")
    assert first.principal_allocated["A-H"] == D("411440223.85")
    assert first.principal_allocated["M-1"] == D("14657659.91")
    assert first.principal_allocated["M-1H"] == D("773508.04")
    for tranche in ("M-2A", "M-2AH", "M-2B", "M-2BH", "B-1H", "B-2H", "B-3H"):
        assert first.principal_allocated[tranche] == D(0)
    assert all(v == D(0) for v in first.write_down_allocated.values())
    assert sum(first.principal_allocated.values(), D(0)) == first.stated_principal


def test_stress_write_downs_hit_first_loss_first(stress: RunResult) -> None:
    first = stress.structure[0]
    assert first.tranche_write_down > D(0)
    assert first.write_down_allocated["B-3H"] == first.tranche_write_down
    assert stress.tranche_summary_for("B-3H").total_write_downs > D(0)
    # Cumulative write-downs are non-decreasing and end at the lifetime total.
    for tranche in TRANCHE_ORDER:
        running = D(0)
        for row in stress.structure:
            running += row.write_down_allocated[tranche]
            assert row.cumulative_write_down[tranche] == running


def test_tranche_allocations_csv(pricing: RunResult, tmp_path: Path) -> None:
    written = {p.name: p for p in export_csv(pricing, tmp_path)}
    with written["tranche_allocations.csv"].open(encoding="utf-8", newline="") as handle:
        rows = list(csv.DictReader(handle))
    assert list(rows[0]) == [
        "payment_date_number",
        "payment_date",
        "tranche",
        "balance_before",
        "principal",
        "write_down",
        "write_up",
        "increase",
        "cumulative_write_down",
        "balance_after",
    ]
    assert len(rows) == len(pricing.structure) * len(TRANCHE_ORDER)
    assert rows[1]["tranche"] == "A-1" and rows[1]["principal"] == "10346250.00"
    assert "tranche_summary.csv" in written


# --------------------------------------------------------------------------------------
# Workbook
# --------------------------------------------------------------------------------------


def _sheet_values(path: Path) -> dict[str, list[list[Any]]]:
    workbook = load_workbook(path)
    return {name: [list(row) for row in workbook[name].iter_rows(values_only=True)] for name in workbook.sheetnames}


def test_export_structure_workbook(pricing: RunResult, tmp_path: Path) -> None:
    path = tmp_path / "pricing.xlsx"
    export_structure_workbook(pricing, path)
    workbook = load_workbook(path)
    assert tuple(workbook.sheetnames) == SHEET_NAMES

    readme = [row for row in workbook["README"].iter_rows(values_only=True)]
    assert any(VALUES_ONLY_SENTENCE in str(cell) for row in readme for cell in row)
    assert any(pricing.manifest.run_id == cell for row in readme for cell in row)

    structure = workbook["Structure"]
    headers = [cell.value for cell in structure[1]]
    assert headers == list(STRUCTURE_HEADERS)
    assert headers[-12:] == [f"{t} after" for t in BOTTOM_UP_ORDER]
    first = dict(zip(headers, [cell.value for cell in structure[2]], strict=True))
    assert isinstance(first["Payment Date"], datetime)
    assert first["Payment Date"].date() == date(2026, 3, 25)
    assert first["Senior %"] == float(D("0.9647500"))
    assert first["Senior Reduction Amount"] == float(D("422332461.93"))
    # Spec 02 section 12: Class Notional Amounts after Payment Date 1.
    expected_after = {
        "A-H": "21276215056.99",
        "A-1": "265553750.00",
        "A-1H": "14013693.92",
        "M-1": "261242340.09",
        "M-1H": "13786173.96",
        "M-2A": "37850000.00",
        "M-2AH": "2017015.00",
        "M-2B": "37850000.00",
        "M-2BH": "2017015.00",
        "B-1H": "102515181.00",
        "B-2H": "273373818.00",
        "B-3H": "56953878.00",
    }
    for tranche, text in expected_after.items():
        assert first[f"{tranche} after"] == float(D(text)), tranche
    assert structure.cell(row=2, column=headers.index("A-H after") + 1).number_format == "#,##0.00"
    assert structure.cell(row=2, column=headers.index("Senior %") + 1).number_format == "0.00000%"
    assert structure.freeze_panes == "C2"
    assert structure.max_row == len(pricing.structure) + 1

    principal = workbook["Tranche principal"]
    assert [c.value for c in principal[1]][2:] == list(BOTTOM_UP_ORDER)
    assert principal.cell(row=2, column=3 + BOTTOM_UP_ORDER.index("A-1")).value == float(D("10346250.00"))

    notes = workbook["Notes"]
    assert notes["C1"].value == "Class A-1" and notes["C2"].value == "Beginning balance"
    assert notes["E3"].value == float(D("10346250.00"))  # A-1 principal on Payment Date 1
    assert notes["D3"].value == float(D("1243718.57"))  # A-1 interest on Payment Date 1

    assert len(workbook["Charts"]._charts) == EXPECTED_CHART_COUNT
    assert workbook.properties.created == datetime(2026, 1, 1, tzinfo=workbook.properties.created.tzinfo)


def test_workbook_is_deterministic(pricing: RunResult, project_root: Path, tmp_path: Path) -> None:
    assert structure_workbook_bytes(pricing) == structure_workbook_bytes(pricing)
    other = _run(project_root, "pricing_speed.yaml", "1970-01-01T00:00:00Z")
    assert other.manifest.run_id == pricing.manifest.run_id
    first_path = tmp_path / "first.xlsx"
    second_path = tmp_path / "second.xlsx"
    export_structure_workbook(pricing, first_path)
    export_structure_workbook(other, second_path)
    first_values = _sheet_values(first_path)
    second_values = _sheet_values(second_path)
    assert set(first_values) == set(second_values) == set(SHEET_NAMES)
    differences = [
        (sheet, row_index, cell_index, a, b)
        for sheet in SHEET_NAMES
        for row_index, (row_a, row_b) in enumerate(
            zip(first_values[sheet], second_values[sheet], strict=True), start=1
        )
        for cell_index, (a, b) in enumerate(zip(row_a, row_b, strict=True), start=1)
        if a != b
    ]
    assert [(d[0], d[3], d[4]) for d in differences] == [
        ("README", "2026-01-01T00:00:00Z", "1970-01-01T00:00:00Z")
    ]


def test_stress_workbook_exports(stress: RunResult, tmp_path: Path) -> None:
    path = tmp_path / "stress.xlsx"
    export_structure_workbook(stress, path)
    workbook = load_workbook(path)
    write_downs = workbook["Tranche write-downs"]
    assert write_downs.max_row == 241
    b3h = 3 + BOTTOM_UP_ORDER.index("B-3H")
    assert write_downs.cell(row=2, column=b3h).value == float(stress.structure[0].tranche_write_down)
    summary = workbook["Summary"]
    assert summary["A2"].value == "Note"
    assert isinstance(summary["A3"].value, str) and summary["A3"].value == "A-1"


def test_cli_writes_workbook_with_overrides(project_root: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    out = tmp_path / "cli.xlsx"
    code = excel_main(
        [
            str(project_root / "scenarios/pricing_speed.yaml"),
            str(out),
            "--cer",
            "0.5",
            "--early",
            "off",
            "--project-root",
            str(project_root),
        ]
    )
    assert code == 0 and out.is_file()
    assert "wrote" in capsys.readouterr().out
    workbook = load_workbook(out)
    readme = {row[0]: row[1] for row in workbook["README"].iter_rows(min_row=3, max_row=14, values_only=True)}
    assert readme["Scenario"] == "pricing-speed+overrides"
    assert "CER 0.50%" in str(readme["Scenario label"])
    assert "to scheduled maturity" in str(readme["Scenario label"])
    assert readme["Scenario source"] == "(built from values, no file)"
    assert workbook["Structure"].max_row == 241
    assert workbook["Inputs"]["A1"].value == "Item"
    assert date.fromisoformat("2026-03-25") == workbook["Structure"]["B2"].value.date()


def test_cli_rejects_bad_override(project_root: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    code = excel_main(
        [str(project_root / "scenarios/pricing_speed.yaml"), str(tmp_path / "x.xlsx"), "--cpr", "150"]
    )
    assert code == 1
    assert "CPR 1.50 must satisfy 0 <= CPR < 1" in capsys.readouterr().err
    assert not (tmp_path / "x.xlsx").exists()


def test_charts_stack_in_presentation_order_with_presentation_colours(pricing: RunResult) -> None:
    """The workbook's bottom-up order is the shared one (each Note directly under its H
    tranche) and every stacked series carries ``crt.presentation``'s colour."""
    from crt.presentation import STACK_ORDER_BOTTOM_UP, tranche_colour

    assert BOTTOM_UP_ORDER == STACK_ORDER_BOTTOM_UP
    for note in NOTE_CLASSES:
        assert BOTTOM_UP_ORDER[BOTTOM_UP_ORDER.index(note) + 1] == note + "H"
    workbook = build_structure_workbook(pricing)
    charts = workbook["Charts"]._charts
    assert len(charts) == EXPECTED_CHART_COUNT
    stack_all, stack_ex, write_downs, cumulative = charts[:4]
    for chart, tranches in (
        (stack_all, BOTTOM_UP_ORDER),
        (stack_ex, tuple(t for t in BOTTOM_UP_ORDER if t != "A-H")),
        (write_downs, BOTTOM_UP_ORDER),
        (cumulative, BOTTOM_UP_ORDER),
    ):
        assert chart.grouping == "stacked"
        fills = [s.graphicalProperties.solidFill.srgbClr for s in chart.series]
        colours = [getattr(fill, "val", fill) for fill in fills]  # str, or RGB(.val)
        assert colours == [tranche_colour(t) for t in tranches]
    # README legend: top of the stack first, mirroring the charts.
    readme = list(workbook["README"].iter_rows(values_only=True))
    legend_start = next(i for i, row in enumerate(readme) if row[0] == "Tranche") + 1
    legend = [row[0] for row in readme[legend_start : legend_start + len(BOTTOM_UP_ORDER)]]
    assert legend == list(reversed(BOTTOM_UP_ORDER))
